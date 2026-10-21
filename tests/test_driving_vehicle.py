"""Tests for title-based vehicle recognition, the relaxed sheet admission, dim_vehicle.csv and the two-relation distance fix."""
import datetime as dt
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pipeline.edit_driving import EditConfig, edit  # noqa: E402
from pipeline.make_dummy import gen_clean, inject_man_distance  # noqa: E402
from pipeline.parse_driving import (  # noqa: E402
    assign_vehicle_ids, load_vehicle_table, model_key_of, parse_folder, split_periods, to_frame,
)


def _book(path, title, header, rows, header_row=3):
    wb = Workbook()
    ws = wb.active
    if title:
        ws.cell(1, 1, title)
    for c, h in enumerate(header, 1):
        ws.cell(header_row, c, h)
    for i, r in enumerate(rows):
        for c, v in enumerate(r, 1):
            ws.cell(header_row + 1 + i, c, v)
    wb.save(path)


def _days(n):
    return [dt.datetime(2024, 1, 1) + dt.timedelta(days=i) for i in range(n)]


def _parse(folder):
    files, rows, recon, *_ = parse_folder(folder)
    return to_frame(rows), recon


def test_model_words():
    assert model_key_of("대우프리마 트렉터 일일 운행기록") == "Daewoo"
    assert model_key_of("스카니아 트렉터 일일 운행기록") == "Scania"
    assert model_key_of("MAN TGX 운행일지") == "MAN"
    assert model_key_of("대우 신형 운행기록") == ""  # plain Daewoo is not Prima
    assert model_key_of("MANAGER 운행") == ""


def test_title_beats_structure_and_new_car_is_loaded(tmp_path):
    # a Daewoo title on a sheet shaped like another layout: the title wins
    _book(tmp_path / "a.xlsx", "대우프리마 트렉터 일일 운행기록", ["날짜", "1일 주행거리", "연비", "평균 운행속도"],
          [[d, 200.0 + i, 3.1, 50.0] for i, d in enumerate(_days(5))])
    # a new car: unknown model word, a column nobody knows
    _book(tmp_path / "b.xlsx", "신형 트렉터 일일 운행기록", ["날짜", "1일 주행거리", "엔진온도"],
          [[d, 100.0 + i, 90 + i] for i, d in enumerate(_days(4))])
    df, recon = _parse(tmp_path)
    assert len(df) == 9
    assert all(r["balanced"] for r in recon)
    a = df[df["vehicle"] == "Daewoo"]
    assert len(a) == 5 and set(a["vehicle_source"]) == {"title"} and set(a["vehicle_model_key"]) == {"Daewoo"}
    b = df[df["vehicle"].str.startswith("unk-")]
    assert len(b) == 4 and set(b["vehicle_source"]) == {"title_unknown_model"} and set(b["vehicle_model_key"]) == {""}
    assert b["vehicle_title"].iloc[0].startswith("신형")
    assert "엔진온도" in b["extra"].iloc[0]
    assert (b["vehicle_id"] == "").all()


def test_structure_guess_only_without_title_and_counted(tmp_path):
    _book(tmp_path / "c.xlsx", None, ["날짜", "1일 주행거리", "누적 주행거리", "연비"], [[d, 200.0, 5000.0 + 200 * i, 3.0] for i, d in enumerate(_days(5))], header_row=1)
    df, _ = _parse(tmp_path)
    assert len(df) == 5 and set(df["vehicle_source"]) == {"structure"} and set(df["vehicle"]) == {"Daewoo"}


def test_sales_like_sheet_is_not_a_driving_log(tmp_path):
    # date column, client-name columns with numbers, no driving title, no distance column
    _book(tmp_path / "s.xlsx", "월별 매출 현황", ["일자", "가나상사", "다라물산"], [[d, 1000 + i, 2000 + i] for i, d in enumerate(_days(6))])
    df, recon = _parse(tmp_path)
    assert df.empty
    assert recon[0]["status"] == "unknown_layout" and recon[0]["quarantined"] == recon[0]["source_rows"]
    assert all(r["balanced"] for r in recon)


def test_date_and_measure_without_known_layout_loads_when_title_says_driving(tmp_path):
    _book(tmp_path / "d.xlsx", "특수차 운행기록", ["날짜", "연료주입량"], [[d, 50.0] for d in _days(3)])
    _book(tmp_path / "e.xlsx", "특수차 정비", ["날짜", "연료주입량"], [[d, 50.0] for d in _days(3)])
    df, recon = _parse(tmp_path)
    assert len(df) == 3
    assert sorted(r["status"] for r in recon) == ["ok", "unknown_layout"]


def test_vehicle_table_assigns_id_only_when_exactly_one_entry_matches(tmp_path):
    csv_path = tmp_path / "v.csv"
    csv_path.write_text(
        "vehicle_id,brand,model,model_year,color,service_start,service_end,title_words,note\n"
        ",,,,,,,,\n"
        "V03,대우,프리마,2019,,2019-01-01,2023-12-31,대우프리마,\n"
        "V01,스카니아,R450,2016,,2016-01-01,2018-12-31,,\n"
        "V04,스카니아,R500,2023,,2024-01-01,,,\n", encoding="utf-8")
    table = load_vehicle_table(csv_path)
    assert [e["id"] for e in table] == ["V03", "V01", "V04"]  # the empty row is ignored
    rows = pd.DataFrame({"vehicle": ["Daewoo", "Daewoo", "Scania", "Scania", "Scania"],
                         "vehicle_title": ["대우프리마 운행", "대우프리마 운행"] + ["스카니아 운행"] * 3,
                         "vehicle_model_key": ["Daewoo", "Daewoo", "Scania", "Scania", "Scania"], "vehicle_id": [""] * 5,
                         "date": pd.to_datetime(["2020-05-01", "2024-05-01", "2017-05-01", "2024-05-01", "2020-05-01"])})
    out, counts = assign_vehicle_ids(rows, table)
    # two Scania are told apart by dates; a date in neither service period gets no id
    assert list(out["vehicle_id"]) == ["V03", "", "V01", "V04", ""]
    assert list(out["vehicle"][[0, 2, 3]]) == ["V03", "V01", "V04"]
    assert counts["assigned"] == 3 and counts["unmatched"] == 2


def test_title_without_model_word_is_matched_by_date_alone(tmp_path):
    csv_path = tmp_path / "v.csv"
    csv_path.write_text(
        "vehicle_id,brand,model,model_year,color,service_start,service_end,title_words,note\n"
        "V02,MAN,,,,2015-01-01,2019-03-31,,\n"
        "V03,Daewoo,,,,2019-04-01,2023-08-31,,\n", encoding="utf-8")
    rows = pd.DataFrame({"vehicle": ["unk-a1", "unk-a1", "Daewoo"], "vehicle_title": ["트렉터 운행일지"] * 2 + ["대우프리마 운행"],
                         "vehicle_model_key": ["", "", "Daewoo"], "vehicle_id": [""] * 3,
                         "date": pd.to_datetime(["2017-05-01", "2020-05-01", "2018-05-01"])})
    out, counts = assign_vehicle_ids(rows, load_vehicle_table(csv_path))
    # no model word: the date decides (V02 in 2017, V03 in 2020); a known model word outside its period stays unmatched
    assert list(out["vehicle_id"]) == ["V02", "V03", ""]


def test_shipped_vehicle_table_has_four_cars():
    table = load_vehicle_table(Path(__file__).resolve().parent.parent / "pipeline" / "dim_vehicle.csv")
    assert [e["id"] for e in table] == ["V01", "V02", "V03", "V04"]
    assert [e["brand_key"].lower() for e in table] == ["scania", "man", "daewoo", "scania"]


def test_provisional_period_split_for_one_brand_with_another_car_between():
    days = lambda start, n: [pd.Timestamp(start) + pd.Timedelta(days=i) for i in range(n)]  # noqa: E731
    dates = days("2016-01-01", 40) + days("2017-01-01", 40) + days("2018-01-01", 40) + [pd.Timestamp("2018-03-01")]  # last: a stray day
    keys = ["Scania"] * 40 + ["MAN"] * 40 + ["Scania"] * 40 + ["Scania"]
    df = pd.DataFrame({"vehicle": keys, "date": dates, "vehicle_id": [""] * len(keys)})
    out, info = split_periods(df)
    assert info == {"Scania": 2}
    assert list(out["vehicle"].iloc[:40].unique()) == ["Scania-1"] and set(out["vehicle"].iloc[80:]) == {"Scania-2"}
    assert set(out["vehicle"].iloc[40:80]) == {"MAN"}  # a key with one stretch keeps its name
    # a gap with no other car inside is not a break, and a short stretch is a stray row, not a period
    one = pd.DataFrame({"vehicle": ["Scania"] * 60, "date": days("2016-01-01", 30) + days("2017-01-01", 30), "vehicle_id": [""] * 60})
    assert split_periods(one)[1] == {}
    stray = pd.DataFrame({"vehicle": ["Scania"] + ["MAN"] * 40 + ["Scania"] * 40,
                          "date": [pd.Timestamp("2015-01-01")] + days("2016-01-01", 40) + days("2017-01-01", 40), "vehicle_id": [""] * 81})
    assert split_periods(stray)[1] == {}


@pytest.fixture(scope="module")
def dummy():
    return gen_clean(20261018, "dev")


def test_two_relation_fix_only_on_rows_without_chain_check(dummy):
    dirty, truth = inject_man_distance(dummy, 99)
    fixed, corr, _, info = edit(dirty, EditConfig(use_mincell=False))
    cl = dummy.set_index("row_id")
    t_rows = {int(r) for r in truth["row_id"]}
    new = corr[corr["rule"] == "distance_two_relations"]
    assert len(new) > 0
    assert set(new["column"]) == {"distance"}
    # a fix lands within the display unit of the clean value and never touches a sound row
    assert all(int(r.row_id) in t_rows for r in new.itertuples())
    assert all(abs(r.new_value - cl.at[int(r.row_id), "distance"]) <= 1.0 + 1e-9 for r in new.itertuples())
    # rows that the cumulative-distance chain covers never go through the new rule
    chain_rows = set(dummy[dummy["cumulative_distance"].notna()]["row_id"])
    assert not (set(int(x) for x in new["row_id"]) & chain_rows)


def test_second_round_rules_unchanged_by_the_new_rule(dummy):
    dirty, _ = inject_man_distance(dummy, 5)
    _, c_leg, _, _ = edit(dirty, EditConfig(policy="legacy", use_mincell=False))
    _, c_new, _, _ = edit(dirty, EditConfig(use_mincell=False))
    old_rules = c_new[c_new["rule"] != "distance_two_relations"]
    key = lambda d: sorted(zip(d["row_id"], d["column"], d["rule"], d["new_value"]))  # noqa: E731
    assert key(old_rules) == key(c_leg)


def test_warning_instead_of_fix_when_relations_disagree(dummy):
    dirty = dummy.copy().reset_index(drop=True)
    man = dirty[dirty["cumulative_distance"].isna() & (dirty["distance"] > 0) & dirty["consumed_fuel"].notna()]
    i = man.index[len(man) // 2]
    dirty.at[i, "distance"] = float(dirty.at[i, "distance"]) + 80.0
    dirty.at[i, "speed"] = float(dirty.at[i, "speed"]) * 1.7  # now speed x time points elsewhere than fuel efficiency x fuel
    _, corr, _, info = edit(dirty, EditConfig(use_mincell=False))
    rid = int(dirty.at[i, "row_id"])
    assert rid not in set(int(x) for x in corr["row_id"])
    assert info["fix_status"][rid] == "relations_disagree"
    assert any(w["row_id"] == rid for w in info["soft_warnings"])
