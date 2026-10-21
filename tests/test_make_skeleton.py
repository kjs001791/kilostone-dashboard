import datetime as dt
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pipeline import make_skeleton as ms  # noqa: E402

# values that must never reach the skeleton
SECRETS = ["52000", "18340", "2017", "4321"]


@pytest.fixture()
def made(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2017년 3월"
    ws["A1"] = "2017년 3월 운송 내역"
    ws.merge_cells("A1:C1")
    ws["A2"], ws["B2"], ws["C2"] = "날짜", "쌍용", "한일"
    for r in range(3, 8):
        ws.cell(r, 1, dt.date(2017, 3, r))
        ws.cell(r, 2, 18340)
        ws.cell(r, 3, 4321.5).number_format = "#,##0"
    ws["A8"] = "합 계"
    ws["B8"] = "=SUM(B3:B7)/1000"
    ws["A9"] = "단가(발행단가)"
    ws["B9"] = 52000
    ws["B10"] = "=B8*B9*1.1+52000"
    ws["C10"] = "='2017년 3월'!B9&\"12일\""
    ws["B9"].comment = Comment("2017 인상분 52000", "x")
    src = tmp_path / "2017년 3월.xlsx"
    wb.save(src)
    out = tmp_path / "out"
    ms.main([str(src), "--out", str(out)])
    return out


def test_no_value_leaks(made):
    text = (made / "skeleton_01.txt").read_text(encoding="utf-8")
    wb = load_workbook(made / "skeleton_01.xlsx")
    cells = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
    comments = " ".join(c.comment.text for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.comment)
    titles = " ".join(ws.title for ws in wb.worksheets)
    for s in SECRETS:
        assert s not in text
        assert s not in cells + comments + titles


def test_layout_and_logic_kept(made):
    text = (made / "skeleton_01.txt").read_text(encoding="utf-8")
    assert "merged: A1:C1" in text
    assert "R2: A=날짜 | B=쌍용 | C=한일" in text
    assert "R3: A=#D{yyyy-mm-dd} | B=#N | C=#N{#,##0}" in text
    assert "R4..R7: same layout as R3 (4 rows)" in text
    assert "F: =SUM(B3:B7)/1000" in text
    assert "F: =B8*B9*1.1+#K" in text
    assert "'9999년 9월'!B9" in text and '"99일"' in text


def test_formula_cells_hold_text_not_formulas(made):
    wb = load_workbook(made / "skeleton_01.xlsx")
    ws = wb.worksheets[0]
    assert ws["B8"].data_type == "s" and ws["B8"].value.startswith("F: =")


def test_formatted_far_cells_do_not_blow_up(tmp_path):
    # a styled empty cell at the last column and a far row stretches the used range to billions of cells
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "거래처"
    ws["XFD1"].number_format = "#,##0"
    ws["A1048576"].number_format = "#,##0"
    src = tmp_path / "wide.xlsx"
    wb.save(src)
    ms.main([str(src), "--out", str(tmp_path / "out")])
    text = (tmp_path / "out" / "skeleton_01.txt").read_text(encoding="utf-8")
    assert "R1: A=거래처" in text


def test_xls_is_refused(tmp_path, capsys):
    ms.main([str(tmp_path / "old.xls"), "--out", str(tmp_path)])
    assert ".xlsx/.xlsm only" in capsys.readouterr().out
