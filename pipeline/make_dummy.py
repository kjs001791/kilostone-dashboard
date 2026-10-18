"""Dummy driving-log generator: clean standard table, injected errors, truth list, and fake workbooks.

The clean table is generated first. Errors are injected afterwards so that the truth is known exactly.
Dev and test sets differ in seed, series length, units, gap rate, error rate, error mix and spacing.
The test set also carries one error kind the dev set never shows (digit_add).
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pipeline.common import KOREAN, ROOT, STD_COLUMNS, load_config, path_from, write_csv  # noqa: E402

PERIODS = [("MAN", "2019-03-31"), ("Daewoo", "2023-11-30"), ("Scania", "2099-12-31")]

SETTINGS = {
    "dev": {
        "start": "2016-01-01", "end": "2026-03-31", "gap_rate": 0.03, "idle_row_rate": 0.7,
        "uc": {"MAN": None, "Daewoo": 1.0, "Scania": 1.0}, "dd": {"MAN": 1, "Daewoo": 1, "Scania": 1},
        "error_rate": 0.12, "spacing": 4, "copy_full_row": 0.5,
        "mix": {"digit_drop": 3, "adjacent_key": 3, "transpose": 2, "copy_previous": 2, "date_abnormal": 1, "two_cells": 2},
    },
    "test": {
        "start": "2016-02-01", "end": "2026-02-28", "gap_rate": 0.07, "idle_row_rate": 0.5,
        "uc": {"MAN": None, "Daewoo": 1.0, "Scania": 0.1}, "dd": {"MAN": 0, "Daewoo": 0, "Scania": 1},
        "error_rate": 0.15, "spacing": 3, "copy_full_row": 0.5,
        "mix": {"digit_drop": 2, "adjacent_key": 2, "transpose": 2, "copy_previous": 2, "date_abnormal": 1, "two_cells": 2, "digit_add": 2},
    },
}
MEAN_D = {"MAN": (230, 70), "Daewoo": (250, 80), "Scania": (270, 80)}
START_ODO = {"MAN": 200000.0, "Daewoo": 40000.0, "Scania": 380000.0}


def vehicle_on(day: pd.Timestamp) -> str:
    for name, last in PERIODS:
        if day <= pd.Timestamp(last):
            return name
    return "Scania"


def gen_clean(seed: int, kind: str) -> pd.DataFrame:
    s = SETTINGS[kind]
    rng = np.random.default_rng(seed)
    odo = dict(START_ODO)
    rows = []
    for day in pd.date_range(s["start"], s["end"]):
        veh = vehicle_on(day)
        wd = day.weekday()
        p_work = 0.88 if wd < 5 else (0.5 if wd == 5 else 0.15)
        working = rng.random() < p_work
        true_d = max(5.0, rng.normal(*MEAN_D[veh])) if working else 0.0
        odo[veh] += true_d
        if rng.random() < s["gap_rate"]:
            continue
        if not working and rng.random() > s["idle_row_rate"]:
            continue
        dec = s["dd"][veh]
        D = round(true_d, dec)
        uc = s["uc"][veh]
        C = np.nan if uc is None else float(np.floor(odo[veh] / uc) * uc)
        if uc is not None:
            C = round(C, 3)
        rec = {"vehicle": veh, "date": day, "distance": D, "cumulative_distance": C}
        if working:
            mins = int(round(D / rng.uniform(42, 58) * 60))
            time_h = mins / 60
            fuel = round(D / rng.uniform(2.6, 3.4), 1)
            rec.update({"time_h": time_h, "speed": round(D / time_h, 1) if time_h else np.nan,
                        "consumed_fuel": fuel, "fuel_efficiency": round(D / fuel, 2) if fuel else np.nan,
                        "refuel": float(rng.integers(100, 400)) if rng.random() < 0.25 else np.nan,
                        "reurea": float(rng.integers(10, 40)) if rng.random() < 0.05 else np.nan})
            if veh == "Scania":
                idle = round(rng.uniform(0.1, 2.0), 1)
                pto = round(rng.uniform(0.0, 1.5), 1)
                rec.update({"fuel_rate_per_hour": round(fuel / max(time_h, 0.1), 1), "consumed_fuel_idle": round(idle * 1.5, 1),
                            "consumed_fuel_pto": round(pto * 4, 1), "time_idle_h": round(idle * 60) / 60, "time_pto_h": round(pto * 60) / 60})
        if veh == "Daewoo":  # this layout has no speed or time columns
            rec["time_h"] = np.nan
            rec["speed"] = np.nan
        rows.append(rec)
    df = pd.DataFrame(rows)
    for c in STD_COLUMNS:
        if c not in df.columns:
            df[c] = np.nan
    df["vehicle"] = df["vehicle"]
    month_key = df["date"].dt.year * 12 + df["date"].dt.month
    df["source_file"] = pd.factorize(month_key, sort=True)[0] + 1
    df["source_sheet_no"] = 1
    df["source_row_no"] = df.groupby("source_file").cumcount() + 10
    df["has_values"] = df[["distance", "consumed_fuel", "time_h"]].notna().any(axis=1)
    df = df.sort_values(["date", "source_row_no"], kind="stable").reset_index(drop=True)
    df["row_id"] = range(1, len(df) + 1)
    return df[STD_COLUMNS + ["has_values"]]


# ------------------------------------------------------------------ error injection
NUMPAD = {7: (0, 0), 8: (0, 1), 9: (0, 2), 4: (1, 0), 5: (1, 1), 6: (1, 2), 1: (2, 0), 2: (2, 1), 3: (2, 2), 0: (3, 0.5)}


def adjacent_digits(d: int) -> list:
    r, c = NUMPAD[d]
    return sorted(k for k, (rr, cc) in NUMPAD.items() if k != d and abs(rr - r) <= 1 and abs(cc - c) <= 1.01)


def corrupt_int(n: int, kind: str, rng) -> int | None:
    s = str(n)
    if kind == "digit_drop":
        if len(s) < 3:
            return None
        k = int(rng.integers(0, len(s)))
        t = s[:k] + s[k + 1:]
        return int(t) if t and not t.startswith("0") else None
    if kind == "digit_add":
        k = int(rng.integers(0, len(s) + 1))
        return int(s[:k] + str(int(rng.integers(1 if k == 0 else 0, 10))) + s[k:])
    if kind == "adjacent_key":
        k = int(rng.integers(0, len(s)))
        opts = adjacent_digits(int(s[k]))
        if not opts or (k == 0 and len(s) > 1):
            opts = [o for o in opts if o != 0]
        if not opts:
            return None
        return int(s[:k] + str(int(rng.choice(opts))) + s[k + 1:])
    if kind == "transpose":
        ks = [k for k in range(len(s) - 1) if s[k] != s[k + 1]]
        if not ks:
            return None
        k = int(rng.choice(ks))
        return int(s[:k] + s[k + 1] + s[k] + s[k + 2:])
    return None


def inject(clean: pd.DataFrame, seed: int, kind: str, mix: dict | None = None, error_rate: float | None = None, spacing: int | None = None, units_override: dict | None = None):
    """Return (dirty, truth). truth rows: row_id, column, error_type, clean_value, dirty_value."""
    s = SETTINGS[kind]
    mix = mix or s["mix"]
    rate = s["error_rate"] if error_rate is None else error_rate
    gap = s["spacing"] if spacing is None else spacing
    rng = np.random.default_rng(seed)
    dirty = clean.copy().reset_index(drop=True)
    truth = []
    types = list(mix)
    probs = np.array([mix[t] for t in types], float)
    probs /= probs.sum()
    units = {}
    for veh, g in clean.groupby("vehicle"):
        uc = s["uc"][veh]
        units[veh] = (uc, 10.0 ** -s["dd"][veh])
    if units_override:
        units = dict(units_override)
    n = len(dirty)
    last_err = {v: -10**9 for v in units}
    order = list(range(n))
    for i in order:
        veh = dirty.at[i, "vehicle"]
        uc, ud = units[veh]
        if i == 0 or rng.random() > rate:
            continue
        if i - last_err[veh] < gap or dirty.at[i - 1, "vehicle"] != veh:
            continue
        t = str(rng.choice(types, p=probs))
        if uc is None and t != "date_abnormal":
            continue
        tol = (uc or 0.0) + ud
        if t == "date_abnormal":
            d = pd.Timestamp(dirty.at[i, "date"])
            yr = int(rng.choice([1902, 1900, 2203, 2092]))
            try:
                nd = d.replace(year=yr)
            except ValueError:
                continue
            dirty.at[i, "date"] = nd
            truth.append((int(dirty.at[i, "row_id"]), "date", t, d, nd))
            last_err[veh] = i
            continue
        cols = ["distance", "cumulative_distance"]
        if t == "two_cells":
            pairs = [(c, k) for c, k in ((c, str(rng.choice(["digit_drop", "adjacent_key", "transpose"]))) for c in cols)]
        else:
            pairs = [(str(rng.choice(cols)), t)]
        made = []
        full_copy = False
        for col, k in pairs:
            unit = ud if col == "distance" else uc
            cv = float(clean.at[i, col])
            if np.isnan(cv):
                made = []
                break
            if k == "copy_previous":
                nv = float(clean.at[i - 1, col])
                # a clerk who copies the day above often copies the whole row, not the distance alone
                full_copy = col == "distance" and bool(s.get("copy_full_row")) and rng.random() < s["copy_full_row"]
            else:
                m = int(round(cv / unit))
                cm = corrupt_int(m, k, rng)
                if cm is None:
                    made = []
                    break
                nv = round(cm * unit, 6)
            if np.isnan(nv) or abs(nv - cv) <= tol + 1e-9 or nv < 0:
                made = []
                break
            made.append((col, k, cv, nv))
        if not made:
            continue
        for col, k, cv, nv in made:
            dirty.at[i, col] = nv
            truth.append((int(dirty.at[i, "row_id"]), col, "two_cells" if t == "two_cells" else k, cv, nv))
        if full_copy:
            for c2 in ("consumed_fuel", "time_h", "fuel_efficiency", "speed"):
                dirty.at[i, c2] = clean.at[i - 1, c2]
        last_err[veh] = i
    tr = pd.DataFrame(truth, columns=["row_id", "column", "error_type", "clean_value", "dirty_value"])
    return dirty, tr


def inject_man_distance(clean: pd.DataFrame, seed: int, rate: float = 0.06, spacing: int = 3):
    """Rows without a cumulative distance (MAN in the dummy) have no chain check, so a wrong daily distance is invisible to it. Inject some to score the two relations.

    Returns (dirty, truth) with truth rows row_id, column, error_type, clean_value, dirty_value. Only rows without a cumulative distance and with a working-day distance are touched.
    """
    rng = np.random.default_rng(seed)
    dirty = clean.copy().reset_index(drop=True)
    truth = []
    last = -10**9
    kinds = ["digit_drop", "adjacent_key", "transpose", "digit_add"]
    for i in range(len(dirty)):
        if not np.isnan(dirty.at[i, "cumulative_distance"]) or i - last < spacing or rng.random() > rate:
            continue
        cv = float(dirty.at[i, "distance"])
        if np.isnan(cv) or cv <= 0 or np.isnan(dirty.at[i, "consumed_fuel"]):
            continue
        unit = 0.1 if abs(cv * 10 - round(cv * 10)) < 1e-6 and abs(cv - round(cv)) > 1e-6 else 1.0
        k = str(rng.choice(kinds))
        cm = corrupt_int(int(round(cv / unit)), k, rng)
        if cm is None:
            continue
        nv = round(cm * unit, 6)
        if abs(nv - cv) <= 2 * unit + 1e-9:
            continue
        dirty.at[i, "distance"] = nv
        truth.append((int(dirty.at[i, "row_id"]), "distance", k, cv, nv))
        last = i
    return dirty, pd.DataFrame(truth, columns=["row_id", "column", "error_type", "clean_value", "dirty_value"])


# ------------------------------------------------------------------ fake workbooks (parser test)
def _hm(h):
    if h is None or (isinstance(h, float) and np.isnan(h)):
        return None
    m = int(round(h * 60))
    return f"{m // 60}:{m % 60:02d}"


def write_workbooks(clean: pd.DataFrame, outdir: Path, max_months: int | None = None) -> int:
    from openpyxl import Workbook

    outdir.mkdir(parents=True, exist_ok=True)
    n_files = 0
    for (fno, veh), g in clean.groupby(["source_file", "vehicle"], sort=True):
        if max_months and fno > max_months:
            break
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        def val(x):
            return None if (isinstance(x, float) and np.isnan(x)) else x

        if veh == "MAN":
            ws.cell(1, 1, "MAN TGX 운행일지")
            hdr = ["날짜", "평균 운행속도", "총 운행시간", "1일 주행거리", "연비", "1일 연료소모량", "연료주입량", "요소수 주입량"]
            for c, h in enumerate(hdr, 1):
                ws.cell(3, c, h)
            r = 4
            for _, x in g.iterrows():
                vals = [x["date"].to_pydatetime(), val(x["speed"]), _hm(x["time_h"]), val(x["distance"]), val(x["fuel_efficiency"]),
                        val(x["consumed_fuel"]), val(x["refuel"]), val(x["reurea"])]
                for c, v in enumerate(vals, 1):
                    ws.cell(r, c, v)
                r += 1
            ws.cell(r, 1, "합계")
            ws.cell(r, 4, float(g["distance"].sum()))
        elif veh == "Daewoo":
            ws.cell(1, 1, "대우프리마 운행기록")
            hdr = ["일자", "1일 주행거리", "총 주행거리", "연비", "1일 연료소모량", "연료 주입량", "요소수 주입량"]
            for c, h in enumerate(hdr, 1):
                ws.cell(2, c, h)
            r = 3
            for _, x in g.iterrows():
                vals = [x["date"].to_pydatetime(), val(x["distance"]), val(x["cumulative_distance"]), val(x["fuel_efficiency"]),
                        val(x["consumed_fuel"]), val(x["refuel"]), val(x["reurea"])]
                for c, v in enumerate(vals, 1):
                    ws.cell(r, c, v)
                r += 1
            ws.cell(r, 1, "합계")
        else:
            ws.cell(1, 1, "스카니아 트렉터 일일 운행기록")
            l1 = ["", "운행", "", "연비개요", "", "", "구동시간", "", "", "", "", "", "", ""]
            l2 = ["", "", "", "", "", "", "", "", "", "", "", "", "", ""]
            l3 = ["날짜", "km(주행)", "km(누적)", "l(TOT)", "l(IDL)", "l(PTO)", "h(TOT)", "h(IDL)", "h(PTO)", "l/h", "km/l", "km/h", "l(경유)", "l(요소수)"]
            for c in range(14):
                if l1[c]:
                    ws.cell(2, c + 1, l1[c])
                ws.cell(4, c + 1, l3[c])
            ws.merge_cells("B2:C2")
            ws.merge_cells("D2:F2")
            ws.merge_cells("G2:I2")
            r = 5
            for _, x in g.iterrows():
                vals = [x["date"].to_pydatetime(), val(x["distance"]), val(x["cumulative_distance"]), val(x["consumed_fuel"]),
                        val(x["consumed_fuel_idle"]), val(x["consumed_fuel_pto"]), _hm(x["time_h"]), _hm(x["time_idle_h"]),
                        _hm(x["time_pto_h"]), val(x["fuel_rate_per_hour"]), val(x["fuel_efficiency"]), val(x["speed"]),
                        val(x["refuel"]), val(x["reurea"])]
                for c, v in enumerate(vals, 1):
                    ws.cell(r, c, v)
                r += 1
        wb.save(outdir / f"dummy_{int(fno):03d}_{veh}.xlsx")
        n_files += 1
    return n_files


def main() -> int:
    cfg = load_config()
    fx = path_from(cfg, "fixtures_dir")
    for kind in ("dev", "test"):
        seed = cfg["seeds"][kind]
        clean = gen_clean(seed, kind)
        dirty, truth = inject(clean, seed + 1, kind)
        write_csv(clean, fx / f"driving_dummy_{kind}_clean.csv")
        write_csv(dirty, fx / f"driving_dummy_{kind}_dirty.csv")
        write_csv(truth, fx / f"driving_dummy_{kind}_truth.csv")
        print(f"{kind}: rows {len(clean)}  injected cells {len(truth)}  by type {truth['error_type'].value_counts().to_dict()}")
    clean = gen_clean(cfg["seeds"]["dev"], "dev")
    n = write_workbooks(clean, fx / "xlsx_dummy", max_months=None)
    print("dummy workbooks:", n)
    return 0


if __name__ == "__main__":
    sys.exit(main())
