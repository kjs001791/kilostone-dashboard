"""Turn a workbook into a skeleton: layout, header words, names and formulas stay, values go.

The user runs this on a raw file; Claude reads only the output, so it can learn a new sheet layout
without seeing tons, amounts, prices or dates.

- number cell -> "#N", date/time cell -> "#D" (number format kept next to it in the text dump)
- text cell, comment, sheet name -> every digit becomes 9 (names stay, "2017년 3월" -> "9999년 9월")
- formula -> kept, except number constants not in KEEP_CONST (-> "#K") and digits in text or sheet names
- formulas are written as text "F: =..." so cached results never reach the output

Usage (repo root): venv\\Scripts\\python.exe pipeline\\make_skeleton.py <file.xlsx> [<file.xlsx> ...]
Output: .claude/skeleton/skeleton_NN.xlsx and skeleton_NN.txt (gitignored). Console: paths and counts only.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.tokenizer import Token, Tokenizer

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / ".claude" / "skeleton"
DIGIT = re.compile(r"\d")
CELL_REF = re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")
# constants that carry logic, not data: unit steps (kg -> ton), VAT, small counts and rounding digits
KEEP_CONST = {str(n) for n in range(11)} | {"100", "1000", "10000", "0.1", "1.1"}
MAX_TEXT = 60


def mask_text(s: str) -> str:
    return DIGIT.sub("9", s)


def mask_formula(f: str) -> str:
    try:
        items = Tokenizer(f).items
    except Exception:
        return "=" + mask_text(f.lstrip("="))
    out = []
    for t in items:
        if t.type == Token.OPERAND and t.subtype == Token.NUMBER:
            out.append(t.value if t.value in KEEP_CONST else "#K")
        elif t.type == Token.OPERAND and t.subtype == Token.TEXT:
            out.append(mask_text(t.value))
        elif t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" in t.value:
            sheet, ref = t.value.rsplit("!", 1)
            out.append(mask_text(sheet) + "!" + ref)
        else:
            out.append(t.value)
    return "=" + "".join(out)


def mask_value(v) -> tuple[str, str] | None:
    """(kind, masked text) for one cell value; None for an empty cell."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if hasattr(v, "text"):  # ArrayFormula and friends
        v = v.text
    if isinstance(v, bool):
        return "bool", str(v)
    if isinstance(v, (int, float)):
        return "num", "#N"
    if isinstance(v, (dt.datetime, dt.date, dt.time, dt.timedelta)):
        return "date", "#D"
    s = str(v)
    if s.startswith("="):
        return "formula", "F: " + mask_formula(s)
    return "text", mask_text(s)


def relative_refs(text: str, row: int) -> str:
    """B5*C5 on row 5 and B6*C6 on row 6 both become B[r+0]*C[r+0], so repeated rows collapse."""
    return CELL_REF.sub(lambda m: f"{m.group(1)}[r{int(m.group(2)) - row:+d}]", text)


def shown(kind: str, text: str, fmt: str) -> str:
    if kind in ("num", "date") and fmt and fmt != "General":
        text = f"{text}{{{fmt}}}"
    return text if len(text) <= MAX_TEXT else text[:MAX_TEXT] + "..."


def skeleton(src: Path, out_xlsx: Path, out_txt: Path) -> Counter:
    wb = load_workbook(src, data_only=False)
    kinds: Counter = Counter()
    lines = [f"# skeleton of {mask_text(src.name)}", f"sheets: {len(wb.worksheets)}", ""]
    for sno, ws in enumerate(wb.worksheets, start=1):
        title = mask_text(ws.title)
        lines.append(f"## sheet {sno:02d} '{title}' rows {ws.max_row} cols {ws.max_column} state {ws.sheet_state}")
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            lines.append("merged: " + ", ".join(merged))
        prev_key, run_start, run_len = None, 0, 0
        body: list[str] = []

        def close_run():
            if run_len:
                body.append(f"R{run_start + 1}..R{run_start + run_len}: same layout as R{run_start} ({run_len} rows)")

        # only cells stored in the file: iter_rows() would create every cell of the used range,
        # which runs to millions when a whole row or column carries formatting
        by_row: dict[int, list] = {}
        for (r, _), c in sorted(ws._cells.items()):
            by_row.setdefault(r, []).append(c)
        for r in sorted(by_row):
            row = by_row[r]
            cells = []
            for c in row:
                m = mask_value(c.value)
                if c.comment is not None:
                    c.comment = Comment(mask_text(c.comment.text or ""), "")
                if m is None:
                    continue
                kind, text = m
                kinds[kind] += 1
                c.value = text
                cells.append((c.column_letter, shown(kind, text, c.number_format)))
            if not cells:
                continue
            key = tuple((col, relative_refs(t, r)) for col, t in cells)
            if key == prev_key and r == run_start + run_len + 1:
                run_len += 1
                continue
            close_run()
            body.append(f"R{r}: " + " | ".join(f"{col}={t}" for col, t in cells))
            prev_key, run_start, run_len = key, r, 0
        close_run()
        lines += body + [""]
        ws.title = f"{sno:02d} {title}"[:31]
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    out_txt.write_text("\n".join(lines), encoding="utf-8")
    return kinds


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, default=OUT_DIR)
    args = ap.parse_args(argv)
    for k, src in enumerate(args.files, start=1):
        if src.suffix.lower() not in (".xlsx", ".xlsm"):
            print(f"{k}: .xlsx/.xlsm only. Open it in Excel and save as .xlsx first")
            continue
        out_xlsx, out_txt = args.out / f"skeleton_{k:02d}.xlsx", args.out / f"skeleton_{k:02d}.txt"
        kinds = skeleton(src, out_xlsx, out_txt)
        print(f"{k}: {out_txt}  cells " + ", ".join(f"{n} {kinds[n]}" for n in sorted(kinds)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
